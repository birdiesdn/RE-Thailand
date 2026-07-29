"""Build the DOCX resume from structured content (ATS-friendly, single column)."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Pt, Inches, RGBColor

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
RIGHT_TAB = Inches(7.0)

doc = Document()

section = doc.sections[0]
section.top_margin = Inches(0.5)
section.bottom_margin = Inches(0.5)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.space_before = Pt(0)


def para(space_after=0, space_before=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if align is not None:
        p.alignment = align
    return p


def run(p, text, bold=False, italic=False, size=10, color=None, caps=False):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    if caps:
        r.font.all_caps = True
    return r


def section_heading(text):
    p = para(space_before=9, space_after=2)
    run(p, text, bold=True, size=10.5, color=ACCENT, caps=True)
    pbdr = p.paragraph_format
    pbdr.keep_with_next = True
    # underline rule
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F3A5F")
    borders.append(bottom)
    # w:pBdr must precede these siblings in the CT_PPr schema order.
    pPr.insert_element_before(
        borders,
        "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
        "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
        "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
        "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
        "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
        "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
        "w:pPrChange",
    )


def role_line(left_bold, left_rest, right_text, italic_left=False):
    """One line: bolded left content, right-aligned date via tab stop."""
    p = para(space_before=4, space_after=1)
    p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
    if left_bold:
        run(p, left_bold, bold=True)
    if left_rest:
        run(p, left_rest, italic=italic_left)
    run(p, "\t")
    run(p, right_text, italic=True)
    return p


def bullet(text_parts):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    if isinstance(text_parts, str):
        text_parts = [(text_parts, False)]
    for text, bold in text_parts:
        run(p, text, bold=bold)
    return p


# ---------------------------------------------------------------- header
p = para(align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
run(p, "SUDHINEE “BIRDIE” TRIDHIP", bold=True, size=19, color=ACCENT)

p = para(align=WD_ALIGN_PARAGRAPH.CENTER, space_after=5)
run(
    p,
    "Berkeley, CA  ·  +1 (510) 207-4817  ·  sudhinee.tri@berkeley.edu  ·  "
    "linkedin.com/in/sudhinee-tridhip",
    size=9,
)

p = para(space_after=1)
run(
    p,
    "Architecture-trained developer with two years of end-to-end development experience in "
    "Bangkok and a UC Berkeley MRED+D completed May 2026. Underwrites mixed-use and residential "
    "deals from raw land through entitlement and delivery. Seeking a Development or Acquisitions "
    "Analyst role in California; authorized to work in the U.S. for up to three years under F-1 "
    "STEM OPT.",
    size=9.5,
)

# ---------------------------------------------------------------- education
section_heading("Education")

role_line("University of California, Berkeley", " — Berkeley, CA", "May 2026")
p = para(space_after=1)
run(p, "Master of Real Estate Development + Design (MRED+D)")

role_line("Chulalongkorn University", " — Bangkok, Thailand", "June 2023")
p = para(space_after=1)
run(p, "B.S. Architectural Design (INDA), Second Class Honors, GPA 3.54/4.00")
p = para(space_after=1)
run(
    p,
    "Exchange: École Nationale Supérieure d’Architecture de Versailles, France "
    "— GPA 4.00/4.00, 2022–2023",
    size=9.5,
)

# ---------------------------------------------------------------- experience
section_heading("Professional Experience")

role_line(
    "The Client Project Co., Ltd.",
    " — Bangkok, Thailand",
    "June 2023 – June 2025",
)
p = para(space_after=2)
run(p, "Real Estate Developer & Project Manager", italic=True)

bullet(
    "Underwrote a 6-acre mixed-use urban revitalization project (Ratchada 19): cash-flow models "
    "projecting a 12% IRR against an 8–10% benchmark, and value-add strategies worth $132M in "
    "incremental asset value."
)
bullet(
    "Led master planning and delivery of a 20-acre mixed-use equestrian, K9, and gymnastics "
    "complex in Minburi, coordinating architects and contractors against budget and schedule."
)
bullet(
    "Built the portfolio strategy and design-standard framework for Ally Village, scaling one "
    "community-mall concept across 8 Bangkok assets on a $320K program budget."
)
bullet(
    "Ran market feasibility and demographic analysis for an “Affordable Wellness” "
    "mixed-use concept, expanding the firm’s addressable market 3x and launching a new "
    "product line."
)
role_line(
    "Superdry (Thailand) Co., Ltd.",
    " — Bangkok, Thailand",
    "July – December 2023",
)
p = para(space_after=2)
run(p, "Graphic Designer", italic=True)
bullet(
    "High-volume campaign artwork across social, digital, and pop-up retail; +32% online "
    "engagement."
)

# ---------------------------------------------------------------- projects
section_heading("Real Estate Projects")

role_line(
    "UC Berkeley Capstone & Development Studio",
    " — 15 Marina Blvd and Seawall 321, San Francisco, CA",
    "Feb – May 2026",
)
bullet(
    "Ran feasibility, zoning, and density-bonus analysis on a 2.6-acre Safeway site, using state "
    "streamlining rather than a discretionary rezoning to hold a compliant 8-story midrise and "
    "remove years of entitlement risk."
)
bullet(
    "Underwrote development strategies on return on cost, cap rate, and yield-on-cost."
)

role_line(
    "ULI Hines Student Urban Design Competition",
    " — “The Hearth,” Austin, TX",
    "January 2026",
)
bullet(
    "Directed development strategy and underwriting for a 3-phase, 332,805 SF mixed-use project "
    "with $680M total development value, projecting a 26.45% levered IRR and 1.92x equity "
    "multiple."
)
bullet(
    "Sized a 369-unit program (56% affordable) plus a med-tech incubator, and built the "
    "investment thesis pitched to judges."
)

role_line(
    "Hack-A-House Competition",
    " — “Stack-A-House,” mass-timber ADUs",
    "September 2025",
)
bullet(
    "Co-led product strategy and feasibility for a scalable mass-timber ADU in supply-constrained "
    "markets: 10–25% cost savings, 3–6 months faster than conventional delivery."
)

role_line(
    "Bangkok Land Watch",
    " — independent acquisition pipeline",
    "2026 – Present",
)
bullet(
    "Maintain a live screening tool for 18 Bangkok land parcels — price per sq.wah, parcel "
    "geometry from title surveys, transit access, catchment spending power, competing supply — "
    "plus an investor brief pricing four corridor strategies."
)

# ---------------------------------------------------------------- achievements
section_heading("Achievements & Community")

bullet(
    [
        ("Winner, Horizontal Development", True),
        (" — RE-CU Junior: Start-Up in Real Estate, Bangkok (2023)", False),
    ]
)
bullet(
    [
        ("Project Manager", True),
        (", Community Creative Playground — built playground for 110 students (2022)", False),
    ]
)
bullet(
    [
        ("Volunteer", True),
        (", Habitat for Humanity East Bay/Silicon Valley — 10 homes, Sequoia Grove (2025)", False),
    ]
)

# ---------------------------------------------------------------- skills
section_heading("Skills")

p = para(space_after=1)
run(p, "Development & Finance: ", bold=True)
run(
    p,
    "cash-flow modeling, underwriting, highest-and-best-use analysis, market feasibility, due "
    "diligence, entitlement and zoning strategy, density bonus, leasing strategy, stakeholder "
    "management",
)

p = para(space_after=1)
run(p, "Software: ", bold=True)
run(
    p,
    "Advanced Excel, Argus Enterprise, CoStar, ESRI ArcGIS, U.S. Census/ACS, Rhinoceros 3D, "
    "Grasshopper, Adobe Creative Suite, Blender",
)

p = para(space_after=1)
run(p, "Languages: ", bold=True)
run(p, "Thai (native), English (fluent), Chinese (basic)")

import sys

out = sys.argv[1] if len(sys.argv) > 1 else "Sudhinee_Tridhip_Resume_2026-07.docx"
doc.save(out)
print("wrote", out)
